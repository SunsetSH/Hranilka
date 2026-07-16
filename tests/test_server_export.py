"""Экспорт VPS-серверов во все 4 формата (TXT/CSV/HTML/XLSX), docs/ТЗ_VPS_Серверы.md §5.

По образцу test_fin_export.py, но независимая ветка (сервер — не финансовый
инструмент): по умолчанию пароли пользователей ОС/панелей, приватный
SSH-ключ и passphrase полностью отсутствуют в выводе; с
include_server_secrets=True — все секреты присутствуют; include_servers=False
— серверов нет ни в одном формате; XLSX — отдельная таблица; CSV-инъекция в
имени сервера нейтрализуется; экспорт ветки с одним сервером (ПКМ "Экспорт…").
"""
import csv

import openpyxl

from hranilka.services import export
from hranilka.services.export import Options

OS_PASSWORD = "SuperRootPass1"
SSH_PRIVATE_KEY = "-----BEGIN OPENSSH PRIVATE KEY-----FAKEDATA"
SSH_PASSPHRASE = "correcthorsebattery"
PANEL_PASSWORD = "PanelSecret9"
SECRETS = (OS_PASSWORD, SSH_PRIVATE_KEY, SSH_PASSPHRASE, PANEL_PASSWORD)


def _theme():
    return {"font": "Consolas", "font_size": 14, "text_color": "#000000",
            "main_bg_color": "#F0F0F0", "tree_bg_color": "#FFFFFF"}


def _build_tree(db):
    """Дерево: сервис со свободным сервером и аккаунт, связанный с сервером."""
    sid = db.add_service("Хостинг")
    acc = db.add_account(sid, "Провайдер-аккаунт", login="user")

    server_id = db.add_server(sid, "Мой VPS")
    storage = db.get_server(server_id)
    storage["payload"] = {
        "v": 1, "hosting": "Aeza / промо-тариф", "host": "203.0.113.10",
        "ssh_port": "22", "extra_ips": ["2a01:...", "10.0.0.9"],
        "os_name": "Ubuntu 24.04",
        "location": "Франкфурт", "paid_until": "2030-01-01", "price": "5 USD/мес",
        "notes": "заметка",
        "os_users": [{"login": "root", "password": OS_PASSWORD,
                      "role": "root", "label": ""}],
        "ssh_keys": [{"label": "основной", "key_type": "ed25519",
                      "public_key": "ssh-ed25519 AAAApublic",
                      "private_key": SSH_PRIVATE_KEY,
                      "passphrase": SSH_PASSPHRASE}],
        "panels": [{"panel_type": "3x-ui", "url": "https://203.0.113.10",
                    "login": "admin", "password": PANEL_PASSWORD}],
    }
    db.save_server_with_links(server_id, storage, [acc])
    return db.export_subtree(), server_id


def _read_text(path, fmt):
    enc = "utf-8-sig" if fmt == "csv" else "utf-8"
    with open(path, encoding=enc) as f:
        return f.read()


def _xlsx_text(path):
    wb = openpyxl.load_workbook(path)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    parts.append(str(c.value))
    return "\n".join(parts)


def _all_formats_text(tree, opts, tmp_path):
    result = {}
    for fmt, (fn, ext, _flt) in export.FORMATS.items():
        out = tmp_path / ("exp" + ext)
        fn(tree, opts, str(out))
        result[fmt] = _xlsx_text(out) if fmt == "xlsx" else _read_text(out, fmt)
    return result


# ─── По умолчанию: секреты скрыты, публичные данные и связь остаются ────────

def test_default_hides_secrets(db, tmp_path):
    tree, _sid = _build_tree(db)
    texts = _all_formats_text(tree, Options(theme=_theme(), title="Т"), tmp_path)
    for fmt, text in texts.items():
        for secret in SECRETS:
            assert secret not in text, (fmt, secret)
        assert "Мой VPS" in text, fmt
        assert "203.0.113.10" in text, fmt              # публичный host — не секрет
        assert "ssh-ed25519 AAAApublic" in text, fmt     # публичный ключ — не секрет
        assert "Провайдер-аккаунт" in text, fmt          # связь сервер->аккаунт
        assert "5 USD/мес" in text, fmt                  # новое поле «Стоимость»
        assert "2a01:..." in text, fmt                   # extra_ips (список)
        assert "10.0.0.9" in text, fmt


# ─── include_server_secrets=True: всё раскрыто ───────────────────────────────

def test_include_server_secrets_reveals_all(db, tmp_path):
    tree, _sid = _build_tree(db)
    opts = Options(theme=_theme(), title="Т", include_server_secrets=True)
    texts = _all_formats_text(tree, opts, tmp_path)
    for fmt, text in texts.items():
        for secret in SECRETS:
            assert secret in text, (fmt, secret)


# ─── Панели: port/label убраны из экспорта (УИ §2026-07-15) ─────────────────

def test_panel_export_ignores_legacy_port_and_label_keys(db, tmp_path):
    """Payload с оставшимися port/label в панели (старый формат до УИ
    §2026-07-15) не всплывает в экспорте — SERVER_LIST_SPECS больше не знает
    этих ключей (тот же payload-объект, БД хранит его как есть, raw)."""
    sid = db.add_service("Хостинг")
    server_id = db.add_server(sid, "VPS Legacy")
    storage = db.get_server(server_id)
    storage["payload"] = {
        "v": 1,
        "panels": [{"panel_type": "3x-ui", "url": "https://legacy",
                    "port": "9999", "login": "admin", "password": "pw",
                    "label": "старая метка"}],
    }
    db.save_server(server_id, storage)
    tree = db.export_subtree()

    path = tmp_path / "panels.csv"
    export.export_csv(tree, Options(theme=_theme(), title="Т"), str(path))
    text = _read_text(path, "csv")
    assert "3x-ui" in text
    assert "9999" not in text
    assert "старая метка" not in text


# ─── include_servers=False: серверов нет ни в одном формате ─────────────────

def test_exclude_servers_omits_records(db, tmp_path):
    tree, _sid = _build_tree(db)
    opts = Options(theme=_theme(), title="Т", include_servers=False)
    texts = _all_formats_text(tree, opts, tmp_path)
    for fmt, text in texts.items():
        assert "Мой VPS" not in text, fmt
        assert "203.0.113.10" not in text, fmt
        for secret in SECRETS:
            assert secret not in text, (fmt, secret)
        assert "Провайдер-аккаунт" in text, fmt          # аккаунт остаётся


# ─── CSV-инъекция в имени сервера нейтрализуется ─────────────────────────────

def test_formula_injection_neutralized_csv_xlsx(db, tmp_path):
    sid = db.add_service("S")
    server_id = db.add_server(sid, "=cmd()")             # опасное имя
    tree = db.export_subtree()
    opts = Options(theme=_theme(), title="Т")

    csv_path = tmp_path / "inj.csv"
    export.export_csv(tree, opts, str(csv_path))
    assert "'=cmd()" in _read_text(csv_path, "csv")

    xlsx_path = tmp_path / "inj.xlsx"
    export.export_xlsx(tree, opts, str(xlsx_path))
    wb = openpyxl.load_workbook(xlsx_path)
    cells = [c.value for ws in wb.worksheets
             for row in ws.iter_rows() for c in row if c.value is not None]
    assert "'=cmd()" in cells
    assert server_id is not None


# ─── XLSX: серверы — отдельная таблица со своими заголовками ────────────────

def test_xlsx_server_table_is_separate_with_own_headers(db, tmp_path):
    tree, _sid = _build_tree(db)
    path = tmp_path / "tables.xlsx"
    export.export_xlsx(tree, Options(theme=_theme(), title="Т"), str(path))

    ws = openpyxl.load_workbook(path).active
    assert ws.cell(1, 1).value == "Путь"                 # первая таблица: аккаунты
    account_row = next(r for r in range(2, ws.max_row + 1)
                       if ws.cell(r, 1).value
                       and "Провайдер-аккаунт" in ws.cell(r, 1).value)
    server_header_row = next(r for r in range(account_row + 2, ws.max_row + 1)
                             if ws.cell(r, 1).value == "Путь"
                             and ws.cell(r, 2).value == "Название")
    assert ws.cell(server_header_row - 1, 1).value is None   # строка-отступ
    names = {ws.cell(r, 2).value for r in range(server_header_row + 1, ws.max_row + 1)}
    assert "Мой VPS" in names


def test_csv_server_table_has_own_headers(db, tmp_path):
    tree, _sid = _build_tree(db)
    path = tmp_path / "tables.csv"
    export.export_csv(tree, Options(theme=_theme(), title="Т"), str(path))

    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    account_row = next(i for i, row in enumerate(rows)
                       if row and "Провайдер-аккаунт" in row[0])
    server_header_row = next(i for i, row in enumerate(rows[account_row + 2:], account_row + 2)
                             if row[:2] == ["Путь", "Название"])
    assert rows[server_header_row - 1] == []
    names = {row[1] for row in rows[server_header_row + 1:] if len(row) > 1}
    assert "Мой VPS" in names


# ─── Экспорт ветки с одним сервером (ПКМ по узлу дерева) ─────────────────────

def test_single_server_export_subtree_and_html(db, tmp_path):
    sid = db.add_service("Хостинг")
    server_id = db.add_server(sid, "Одинокий VPS")
    storage = db.get_server(server_id)
    storage["payload"] = {"v": 1, "hosting": "Timeweb", "host": "198.51.100.7"}
    db.save_server(server_id, storage)

    tree = db.export_subtree("server", server_id)
    assert len(tree) == 1 and tree[0]["name"] == "Одинокий VPS"

    out = tmp_path / "server.html"
    export.export_html(tree, Options(theme=_theme(), title="Одинокий VPS"), str(out))
    html = out.read_text(encoding="utf-8")
    assert "Одинокий VPS" in html
    assert "Timeweb" in html
    assert "198.51.100.7" in html
