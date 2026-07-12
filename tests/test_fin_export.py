"""Экспорт финансовых записей во все 4 формата (TXT/CSV/HTML/XLSX).

Проверяется: по умолчанию номер карты маскирован, критичные секреты
(CVV/PIN/seed/passphrase/приватные ключи) полностью отсутствуют в выводе, связь
с аккаунтом присутствует; с include_fin_secrets — секреты и полный номер есть;
include_fin=False — фин-записей нет; formula injection в payload экранируется в
CSV/XLSX; неизвестный item_type не роняет экспорт.
"""
import openpyxl
import csv

from hranilka.services import export
from hranilka.services.export import Options

# Значения-маркеры секретов — их полное отсутствие проверяем по подстроке.
FULL_NUMBER = "4111111111111234"
MASKED_NUMBER = "**** **** **** 1234"
CVV = "987654"
PIN = "43219"
PASSPHRASE = "extraword25"
SEED = "ability able about above absent"
PRIV_KEY = "KzSecretPrivateKeyValue"
SECRETS = (CVV, PIN, PASSPHRASE, SEED, PRIV_KEY)


def _theme():
    return {"font": "Consolas", "font_size": 14, "text_color": "#000000",
            "main_bg_color": "#F0F0F0", "tree_bg_color": "#FFFFFF"}


def _build_tree(db):
    """Дерево: сервис с картой (связана с аккаунтом) и криптокошельком."""
    sid = db.add_service("Банк")
    acc = db.add_account(sid, "МойАккаунт", login="user")

    card = db.add_fin_item(sid, "bank_card", "Зарплатная")
    st = db.load_fin_item(card)
    st["payload"] = {"v": 1, "card_number": FULL_NUMBER, "expiry": "12/29",
                     "cvv": CVV, "pin": PIN, "cardholder": "IVAN IVANOV",
                     "bank_name": "МойБанк"}
    db.save_fin_item_with_links(card, st, [acc])

    wal = db.add_fin_item(sid, "crypto_wallet", "Холодный")
    stw = db.load_fin_item(wal)
    stw["payload"] = {"v": 1, "vendor": "Ledger", "seed_phrase": SEED,
                      "passphrase": PASSPHRASE,
                      "addresses": [{"network": "BTC", "address": "bc1qpublicaddr",
                                     "label": "основной"}],
                      "private_keys": [{"label": "ключ-1", "key": PRIV_KEY}]}
    db.save_fin_item(wal, stw)
    return db.export_subtree()


def _read_text(path, fmt):
    enc = "utf-8-sig" if fmt == "csv" else "utf-8"
    with open(path, encoding=enc) as f:
        return f.read()


def _xlsx_text(path):
    wb = openpyxl.load_workbook(path)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    parts.append(str(c.value))
    return "\n".join(parts)


def _all_formats_text(tree, opts, tmp_path):
    """{fmt: полный текст файла} для txt/csv/html + xlsx (через ячейки)."""
    result = {}
    for fmt, (fn, ext, _flt) in export.FORMATS.items():
        out = tmp_path / ("exp" + ext)
        fn(tree, opts, str(out))
        result[fmt] = _xlsx_text(out) if fmt == "xlsx" else _read_text(out, fmt)
    return result


# ─── По умолчанию: номер маскирован, секреты скрыты ──────────────────────────

def test_default_masks_and_hides_secrets(db, tmp_path):
    tree = _build_tree(db)
    texts = _all_formats_text(tree, Options(theme=_theme(), title="Т"), tmp_path)
    for fmt, text in texts.items():
        assert MASKED_NUMBER in text, fmt
        assert FULL_NUMBER not in text, fmt
        for secret in SECRETS:
            assert secret not in text, (fmt, secret)
        # публичные данные и связь с аккаунтом остаются
        assert "bc1qpublicaddr" in text, fmt           # адрес — не секрет
        assert "МойАккаунт" in text, fmt               # связь запись→аккаунт


# ─── include_fin_secrets: всё раскрыто ───────────────────────────────────────

def test_include_secrets_reveals_all(db, tmp_path):
    tree = _build_tree(db)
    opts = Options(theme=_theme(), title="Т", include_fin_secrets=True)
    texts = _all_formats_text(tree, opts, tmp_path)
    for fmt, text in texts.items():
        assert FULL_NUMBER in text, fmt
        assert MASKED_NUMBER not in text, fmt          # раскрыт — маски нет
        for secret in SECRETS:
            assert secret in text, (fmt, secret)


# ─── include_fin=False: финансовых записей нет ───────────────────────────────

def test_exclude_fin_omits_records(db, tmp_path):
    tree = _build_tree(db)
    opts = Options(theme=_theme(), title="Т", include_fin=False)
    texts = _all_formats_text(tree, opts, tmp_path)
    for fmt, text in texts.items():
        assert "Зарплатная" not in text, fmt
        assert "Холодный" not in text, fmt
        assert MASKED_NUMBER not in text, fmt
        assert "МойАккаунт" in text, fmt               # аккаунт остаётся


# ─── Formula injection в payload (CSV/XLSX) ──────────────────────────────────

def test_formula_injection_neutralized_csv_xlsx(db, tmp_path):
    sid = db.add_service("S")
    iid = db.add_fin_item(sid, "bank_card", "Карта")
    st = db.load_fin_item(iid)
    st["payload"] = {"bank_name": "=cmd()"}            # опасное значение
    db.save_fin_item(iid, st)
    tree = db.export_subtree()
    opts = Options(theme=_theme(), title="Т")

    csv_path = tmp_path / "inj.csv"
    export.export_csv(tree, opts, str(csv_path))
    assert "'=cmd()" in _read_text(csv_path, "csv")

    xlsx_path = tmp_path / "inj.xlsx"
    export.export_xlsx(tree, opts, str(xlsx_path))
    wb = openpyxl.load_workbook(xlsx_path)
    cells = [c.value for ws in wb.worksheets
             for row in ws.iter_rows() for c in row if c.value is not None]
    assert "'=cmd()" in cells                          # экранировано апострофом


def test_xlsx_financial_table_is_below_accounts_with_own_headers(db, tmp_path):
    tree = _build_tree(db)
    path = tmp_path / "tables.xlsx"
    export.export_xlsx(tree, Options(theme=_theme(), title="Т"), str(path))

    ws = openpyxl.load_workbook(path).active
    assert ws.cell(1, 1).value == "Путь"              # первая таблица: аккаунты
    account_row = next(r for r in range(2, ws.max_row + 1)
                       if ws.cell(r, 1).value and "МойАккаунт" in ws.cell(r, 1).value)
    fin_header_row = next(r for r in range(account_row + 2, ws.max_row + 1)
                          if ws.cell(r, 1).value == "Путь"
                          and ws.cell(r, 2).value == "Название")
    assert ws.cell(fin_header_row - 1, 1).value is None  # строка-отступ
    names = {ws.cell(r, 2).value for r in range(fin_header_row + 1, ws.max_row + 1)}
    assert {"Зарплатная", "Холодный"}.issubset(names)


def test_csv_financial_table_has_own_headers(db, tmp_path):
    tree = _build_tree(db)
    path = tmp_path / "tables.csv"
    export.export_csv(tree, Options(theme=_theme(), title="Т"), str(path))

    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f, delimiter=";"))
    account_row = next(i for i, row in enumerate(rows)
                       if row and "МойАккаунт" in row[0])
    fin_header_row = next(i for i, row in enumerate(rows[account_row + 2:], account_row + 2)
                          if row[:2] == ["Путь", "Название"])
    assert rows[fin_header_row - 1] == []
    names = {row[1] for row in rows[fin_header_row + 1:] if len(row) > 1}
    assert {"Зарплатная", "Холодный"}.issubset(names)


# ─── Экспорт одиночной фин-записи (ПКМ по карте) ─────────────────────────────

def test_single_card_export_subtree_and_html(db, tmp_path):
    """export_subtree, стартуя с узла карты, отдаёт ветку из одной записи;
    HTML содержит её данные (путь ПКМ-экспорта карты из дерева)."""
    sid = db.add_service("Банк")
    card = db.add_fin_item(sid, "bank_card", "Зарплатная")
    st = db.load_fin_item(card)
    st["payload"] = {"v": 1, "card_number": FULL_NUMBER, "expiry": "12/29",
                     "bank_name": "МойБанк"}
    db.save_fin_item(card, st)

    tree = db.export_subtree("card", card)
    assert len(tree) == 1 and tree[0]["name"] == "Зарплатная"

    out = tmp_path / "card.html"
    export.export_html(tree, Options(theme=_theme(), title="Зарплатная"), str(out))
    html = out.read_text(encoding="utf-8")
    assert "Зарплатная" in html
    assert "МойБанк" in html
    assert MASKED_NUMBER in html                       # номер маскирован по умолчанию


# ─── Неизвестный item_type не роняет экспорт ─────────────────────────────────

def test_unknown_item_type_no_crash(db, tmp_path):
    sid = db.add_service("S")
    iid = db.add_fin_item(sid, "bank_card", "Странная")
    # Подменяем тип на отсутствующий в реестре (битые данные).
    with db.conn:
        db.cursor.execute(
            "UPDATE fin_items SET item_type = 'alien_type' WHERE id = ?", (iid,))
    tree = db.export_subtree()
    opts = Options(theme=_theme(), title="Т")
    for fmt, (fn, ext, _flt) in export.FORMATS.items():
        out = tmp_path / ("alien" + ext)
        fn(tree, opts, str(out))
        assert out.exists() and out.stat().st_size > 0, fmt
    # Запись неизвестного типа пропущена деревом (итерация 2) — её имени в
    # экспорте нет, но сам экспорт не падает и файлы создаются.
    assert "Странная" not in _read_text(tmp_path / "alien.txt", "txt")
