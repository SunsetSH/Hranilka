"""Экспорт данных Хранилки в человекочитаемые/переносимые форматы.

Форматы: TXT (блокнот), CSV (таблица для Excel/переноса), HTML (оформленный
документ с картинками), PDF (печать/архив). На вход — структура из
Database.export_subtree(): список узлов дерева, где у каждого узла
type=='account' есть ключи 'card' (как load_account) и 'links'.

ВНИМАНИЕ: экспорт сохраняет выбранные данные в ОТКРЫТОМ виде. Включение паролей
и секретов управляется флагами Options.

Модуль не зависит от Qt, кроме export_pdf (импорт Qt изолирован внутри функции).
"""
import base64
import csv
from datetime import datetime
from html import escape as esc
from pathlib import Path


# Префиксы-«иконки» элементов дерева (как в главном окне).
PREFIX = {"folder": "[+] ", "service": "[o] ", "account": "(i) "}

# Группы полей карточки (для галочек «Что включить»).
# BASIC  — вкладки «База» + «Логин и пароль».
# OTHER  — все остальные вкладки (ПД, секр. вопросы, фраза, 2FA, технические).
# Галерея (картинки) выгружается отдельно (include_gallery) — только в HTML/PDF.
GROUP_BASIC = [
    "Название аккаунта", "Адрес сайта (URL)", "Дата создания",
    "Логин", "Пароль", "Пароль сменён", "Сменять пароль каждые",
    "Заметки", "Связанные аккаунты",
]
GROUP_OTHER = [
    "Имя", "Фамилия", "Отчество", "Дата рождения", "Адрес",
    "Секретные вопросы", "Фраза восстановления", "ID устройства",
    "Резервные коды 2FA", "IP адрес", "Браузер", "ОС",
]

CSV_DELIM = ";"  # разделитель для русского Excel


class Options:
    """Параметры экспорта.

    include_basic   — базовые данные, логин и пароль (вкладки «База»/«Логин»).
    include_other   — остальные поля (ПД, секр. вопросы, фраза, 2FA, технические).
    include_gallery — галерея (изображения и описания); только HTML/PDF.
    theme — словарь цветов/шрифта для HTML/PDF."""

    def __init__(self, include_basic=True, include_other=True,
                 include_gallery=True, title="Вся база", theme=None):
        self.include_basic = include_basic
        self.include_other = include_other
        self.include_gallery = include_gallery
        self.title = title
        self.theme = theme or {}

    def _group_enabled(self, label):
        if label in GROUP_BASIC:
            return self.include_basic
        if label in GROUP_OTHER:
            return self.include_other
        return True


# ─── Общий слой: поля карточки ───────────────────────────────────────────────

def _s(v):
    return "" if v is None else str(v)


def _csv_safe(value):
    """Нейтрализует formula injection в табличных форматах (CSV/XLSX).

    Excel/openpyxl трактуют значение, начинающееся с = + - @ (или с управляющего
    символа перед ними), как формулу — вплоть до DDE/WEBSERVICE-вызовов при
    открытии файла. Префиксуем такие значения апострофом: Excel покажет текст
    как есть и формулу не выполнит. На HTML/TXT не влияет (там не вызывается)."""
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        return "'" + value
    return value


def _account_rows(card, links, opts, skip_empty=True):
    """Упорядоченный список (подпись, значение) для одной карточки.

    Поля выключённых групп (basic/other) полностью пропускаются. skip_empty=True
    — пропускать также пустые поля включённых групп (TXT/HTML); False — оставлять
    их пустыми (CSV, фиксированные колонки)."""
    f = card.get("fields", {})
    p = card.get("personal", {})
    rows = []

    def add(label, value):
        if not opts._group_enabled(label):
            return
        value = _s(value)
        if value == "" and skip_empty:
            return
        rows.append((label, value))

    add("Название аккаунта", f.get("account_name"))
    add("Адрес сайта (URL)", f.get("url"))
    add("Дата создания", f.get("creation_date"))
    add("Логин", f.get("login"))
    add("Пароль", f.get("password"))
    add("Пароль сменён", f.get("password_changed_date"))
    iv = f.get("password_change_interval_days")
    add("Сменять пароль каждые", f"{iv} дн." if iv else "")
    add("Заметки", f.get("notes"))
    add("Связанные аккаунты", "\n".join(l["name"] for l in links) if links else "")

    add("Имя", p.get("first_name"))
    add("Фамилия", p.get("last_name"))
    add("Отчество", p.get("middle_name"))
    add("Дата рождения", p.get("birth_date"))
    add("Адрес", p.get("address"))

    q_lines = []
    for q in card.get("questions", []):
        qt = _s(q.get("q")).strip()
        at = _s(q.get("a"))
        if not qt and not at:
            continue
        q_lines.append(f"{qt} — {at}" if qt else at)
    add("Секретные вопросы", "\n".join(q_lines))

    add("Фраза восстановления", card.get("recovery", {}).get("phrase"))
    add("ID устройства", card.get("recovery", {}).get("device_id"))

    codes = [_s(c) for c in card.get("codes", []) if _s(c).strip()]
    add("Резервные коды 2FA", "\n".join(codes))

    add("IP адрес", f.get("ip"))
    add("Браузер", f.get("browser"))
    add("ОС", f.get("os"))

    return rows


# ─── Картинки ────────────────────────────────────────────────────────────────

def _img_ext(data):
    if data[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if data[:4] == b"\x89PNG":
        return ".png"
    if data[:3] == b"GIF":
        return ".gif"
    if data[:2] == b"BM":
        return ".bmp"
    return ".png"


def _img_mime(data):
    return {".jpg": "image/jpeg", ".png": "image/png",
            ".gif": "image/gif", ".bmp": "image/bmp"}[_img_ext(data)]


def _gallery(node):
    """Картинки аккаунта с непустыми данными."""
    return [g for g in node["card"].get("gallery", []) if g.get("data") is not None]


# ─── TXT ─────────────────────────────────────────────────────────────────────

def export_txt(tree, opts, path):
    path = Path(path)
    bar = "=" * 60
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [bar, " ХРАНИЛКА — ЭКСПОРТ", f" {opts.title}    {now}", bar, ""]

    def account(node, indent):
        for label, value in _account_rows(node["card"], node.get("links"), opts):
            if "\n" in value:
                lines.append(f"{indent}{label}:")
                for vl in value.split("\n"):
                    lines.append(f"{indent}    {vl}")
            else:
                lines.append(f"{indent}{label}: {value}")
        # Картинки в текст не помещаются — отмечаем только их наличие
        # (как часть «остальных полей»).
        gal = _gallery(node)
        if gal and opts.include_other:
            lines.append(f"{indent}[Галерея: {len(gal)} изобр. — доступно в HTML/PDF]")
        lines.append(f"{indent}{'-' * 40}")

    def walk(node, depth):
        indent = "    " * depth
        lines.append(f"{indent}{PREFIX[node['type']]}{node['name']}")
        if node["type"] == "account":
            account(node, indent + "    ")
        for child in node.get("children", []):
            walk(child, depth + 1)

    for node in tree:
        walk(node, 0)

    path.write_text("\n".join(lines), encoding="utf-8")


# ─── CSV ─────────────────────────────────────────────────────────────────────

def export_csv(tree, opts, path):
    # Колонки зависят от включённых групп. Колонка «Картинок» (количество
    # вложений) относится к «остальным полям» — выводится только с include_other.
    cols, header = _table_columns(opts)
    out = []

    def walk(node, prefix):
        name = node["name"]
        if node["type"] == "account":
            out.append(_table_row(node, prefix + [name], cols, opts))
        for child in node.get("children", []):
            walk(child, prefix + [name])

    for node in tree:
        walk(node, [])

    # utf-8-sig (BOM) — чтобы Excel корректно открыл кириллицу.
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=CSV_DELIM)
        w.writerow(header)
        w.writerows(out)


def _table_columns(opts):
    """Колонки таблицы (CSV/XLSX) по включённым группам. Возвращает (cols, header),
    где cols — поля карточки, header — полная строка заголовка с «Путь»/«Картинок»."""
    cols = []
    if opts.include_basic:
        cols += GROUP_BASIC
    if opts.include_other:
        cols += GROUP_OTHER
    header = ["Путь"] + cols + (["Картинок"] if opts.include_other else [])
    return cols, header


def _table_row(node, path_parts, cols, opts):
    """Строка таблицы для одного аккаунта (CSV/XLSX)."""
    d = dict(_account_rows(node["card"], node.get("links"), opts, skip_empty=False))
    row = [" / ".join(path_parts)] + [d.get(l, "") for l in cols]
    if opts.include_other:
        row.append(str(len(_gallery(node))))
    # Нейтрализуем формулы в каждой ячейке (CSV и XLSX используют эту строку).
    return [_csv_safe(_s(c)) for c in row]


# ─── HTML ────────────────────────────────────────────────────────────────────

def _default_img_tag(data):
    src = f"data:{_img_mime(data)};base64," + base64.b64encode(data).decode("ascii")
    return f"<img class='shot' src='{src}'>"


def _html_document(tree, opts, render_img=_default_img_tag):
    th = opts.theme
    font = th.get("font", "Consolas")
    fs = int(th.get("font_size", 14))
    text = th.get("text_color", "#000000")
    main_bg = th.get("main_bg_color", "#F0F0F0")
    tree_bg = th.get("tree_bg_color", "#FFFFFF")
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    p = []
    p.append("<!DOCTYPE html><html><head><meta charset='utf-8'>")
    p.append("<title>Хранилка — экспорт</title><style>")
    p.append(f"body{{background:{main_bg};color:{text};"
             f"font-family:'{font}',monospace;font-size:{fs}px;margin:24px;}}")
    p.append(f".banner{{border:2px solid #808080;background:{tree_bg};"
             f"padding:12px 16px;margin-bottom:20px;}}")
    p.append(f".banner h1{{margin:0;font-size:{fs + 8}px;letter-spacing:2px;}}")
    p.append(f"h1.folder{{font-size:{fs + 6}px;border-bottom:2px solid #808080;margin-top:24px;}}")
    p.append(f"h2.service{{font-size:{fs + 3}px;margin-left:12px;}}")
    p.append(f".account{{border:2px inset #808080;background:{tree_bg};"
             f"margin:10px 0 10px 24px;padding:8px 12px;}}")
    p.append(f".account h3{{margin:0 0 6px 0;font-size:{fs + 2}px;}}")
    p.append("table.card{border-collapse:collapse;width:100%;}")
    p.append("table.card td{border:1px solid #808080;padding:4px 8px;vertical-align:top;}")
    p.append("td.label{font-weight:bold;white-space:nowrap;width:200px;}")
    p.append("td.value{white-space:pre-wrap;}")
    p.append(".gallery{margin-top:8px;}")
    p.append("img.shot{max-width:480px;border:2px outset #808080;margin:6px 0;display:block;}")
    p.append(".cap{font-style:italic;margin-bottom:6px;}")
    p.append("</style></head><body>")
    p.append(f"<div class='banner'><h1>ХРАНИЛКА — ЭКСПОРТ</h1>"
             f"{esc(opts.title)} &nbsp;&nbsp; {now}</div>")

    def account(node):
        p.append("<div class='account'>")
        p.append(f"<h3>(i) {esc(node['name'])}</h3><table class='card'>")
        for label, value in _account_rows(node["card"], node.get("links"), opts):
            p.append(f"<tr><td class='label'>{esc(label)}</td>"
                     f"<td class='value'>{esc(value)}</td></tr>")
        p.append("</table>")
        if opts.include_gallery:
            gal = _gallery(node)
            if gal:
                p.append("<div class='gallery'>")
                for g in gal:
                    p.append(render_img(g["data"]))
                    desc = esc(_s(g.get("desc")))
                    if desc:
                        p.append(f"<div class='cap'>{desc}</div>")
                p.append("</div>")
        p.append("</div>")

    def walk(node):
        t = node["type"]
        if t == "folder":
            p.append(f"<h1 class='folder'>[+] {esc(node['name'])}</h1>")
        elif t == "service":
            p.append(f"<h2 class='service'>[o] {esc(node['name'])}</h2>")
        elif t == "account":
            account(node)
        for child in node.get("children", []):
            walk(child)

    for node in tree:
        walk(node)
    p.append("</body></html>")
    return "".join(p)


def export_html(tree, opts, path):
    Path(path).write_text(_html_document(tree, opts), encoding="utf-8")


# ─── PDF ─────────────────────────────────────────────────────────────────────

# Картинки в PDF: ограничиваем размер, чтобы на лист A4 помещалось ~5 штук
# по высоте. Пропорции сохраняются (вписываем в MAXW × MAXH).
_PDF_IMG_MAXW = 460
_PDF_IMG_MAXH = 200


def _pdf_html(tree, opts, render_img):
    """HTML для PDF: компактные заголовки относительно кегля из настроек, без
    фоновых заливок (фон рисуется вручную на всю страницу — см. export_pdf),
    рамки таблицы цветом текста темы."""
    th = opts.theme
    font = th.get("font", "Consolas")
    fs = int(th.get("font_size", 14))
    text = th.get("text_color", "#000000")
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    def val(v):
        return esc(_s(v)).replace("\n", "<br>")

    p = []
    p.append("<html><head><meta charset='utf-8'><style>")
    p.append(f"body{{color:{text};font-family:'{font}',monospace;font-size:{fs}px;}}")
    p.append(f"h1{{font-size:{fs + 2}px;margin:8px 0 2px 0;}}")
    p.append(f"h2{{font-size:{fs + 1}px;margin:6px 0 2px 0;}}")
    p.append(f"h3{{font-size:{fs}px;margin:6px 0 2px 0;}}")
    p.append("table{border-collapse:collapse;width:100%;margin:2px 0 10px 0;}")
    p.append(f"td{{border:1px solid {text};padding:3px 6px;vertical-align:top;}}")
    p.append("td.l{font-weight:bold;}")
    p.append("</style></head><body>")
    p.append(f"<div style='font-size:{fs + 4}px;font-weight:bold;'>ХРАНИЛКА — ЭКСПОРТ</div>")
    p.append(f"<div>{esc(opts.title)} &nbsp; {now}</div>")

    def account(node):
        p.append(f"<h3>(i) {esc(node['name'])}</h3><table>")
        for label, value in _account_rows(node["card"], node.get("links"), opts):
            p.append(f"<tr><td class='l'>{esc(label)}</td><td>{val(value)}</td></tr>")
        p.append("</table>")
        if opts.include_gallery:
            for g in _gallery(node):
                p.append(render_img(g["data"]))
                desc = esc(_s(g.get("desc")))
                if desc:
                    p.append(f"<div><i>{desc}</i></div>")

    def walk(node):
        t = node["type"]
        if t == "folder":
            p.append(f"<h1>[+] {esc(node['name'])}</h1>")
        elif t == "service":
            p.append(f"<h2>[o] {esc(node['name'])}</h2>")
        elif t == "account":
            account(node)
        for child in node.get("children", []):
            walk(child)

    for node in tree:
        walk(node)
    p.append("</body></html>")
    return "".join(p)


def export_pdf(tree, opts, path):
    """PDF из того же HTML, что и export_html (отличная вёрстка), отрисованный
    движком Chromium (QtWebEngine) на страницы A4 — так PDF визуально совпадает
    с HTML-экспортом. Если QtWebEngine недоступен, откатываемся на отрисовку
    через QTextDocument (упрощённая вёрстка)."""
    try:
        _export_pdf_webengine(tree, opts, path)
    except Exception:
        _export_pdf_textdoc(tree, opts, path)


def _export_pdf_webengine(tree, opts, path):
    from PySide6.QtWebEngineCore import QWebEnginePage
    from PySide6.QtCore import QEventLoop, QMarginsF, QUrl, QTimer
    from PySide6.QtGui import QPageLayout, QPageSize
    import tempfile
    import os as _os

    html = _html_document(tree, opts)   # тот же HTML, что и в export_html
    # Грузим из временного файла (обходит лимит setHtml ~2 МБ и корректно
    # подхватывает встроенные картинки data:base64).
    tmp = tempfile.NamedTemporaryFile(
        suffix=".html", delete=False, mode="w", encoding="utf-8")
    tmp.write(html)
    tmp.close()

    page = QWebEnginePage()
    loop = QEventLoop()
    state = {"ok": False, "loaded": False}

    def on_pdf(_fp, ok):
        state["ok"] = ok
        loop.quit()

    def on_load(ok):
        state["loaded"] = ok
        if not ok:
            loop.quit()
            return
        layout = QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4),
            QPageLayout.Orientation.Portrait,
            QMarginsF(10, 10, 10, 10),          # поля в мм
        )
        page.printToPdf(str(path), layout)

    page.loadFinished.connect(on_load)
    page.pdfPrintingFinished.connect(on_pdf)
    page.load(QUrl.fromLocalFile(tmp.name))
    QTimer.singleShot(60000, loop.quit)         # страховочный таймаут
    try:
        loop.exec()
    finally:
        try:
            _os.unlink(tmp.name)
        except OSError:
            pass
    if not state["ok"]:
        raise RuntimeError("QtWebEngine не смог сформировать PDF")


def _export_pdf_textdoc(tree, opts, path):
    # Запасной способ (без QtWebEngine). Импорт Qt изолирован здесь.
    from PySide6.QtGui import (QTextDocument, QPdfWriter, QPageSize, QImage,
                               QFont, QPainter, QColor)
    from PySide6.QtCore import QUrl, QMarginsF, QSizeF, QRectF

    th = opts.theme
    main_bg = th.get("main_bg_color", "#F0F0F0")

    # 96 DPI + нулевые поля страницы: текст нормального кегля (а не крошечного,
    # как при дефолтных 1200 DPI), а фон темы заполняет весь лист. Отступ от краёв
    # даёт documentMargin, поэтому белых полос по краям не остаётся.
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setResolution(96)
    writer.setPageMargins(QMarginsF(0, 0, 0, 0))
    page_w, page_h = writer.width(), writer.height()

    doc = QTextDocument()
    doc.setDefaultFont(QFont(th.get("font", "Consolas"), int(th.get("font_size", 14))))
    doc.setDocumentMargin(24)
    doc.setPageSize(QSizeF(page_w, page_h))

    counter = {"n": 0}

    def render_img(data):
        url = f"mem://img{counter['n']}"
        counter["n"] += 1
        img = QImage()
        img.loadFromData(data)
        doc.addResource(QTextDocument.ResourceType.ImageResource, QUrl(url), img)
        w = img.width() or 1
        h = img.height() or 1
        scale = min(_PDF_IMG_MAXW / w, _PDF_IMG_MAXH / h, 1.0)
        nw, nh = max(1, int(w * scale)), max(1, int(h * scale))
        return f"<img src='{url}' width='{nw}' height='{nh}'>"

    doc.setHtml(_pdf_html(tree, opts, render_img=render_img))

    # Рисуем сами: на каждой странице сначала заливаем фон темы, затем кладём
    # соответствующий срез документа. Это даёт сплошной фон на всех страницах.
    painter = QPainter(writer)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    painter.setRenderHint(QPainter.RenderHint.SmoothPixmapTransform)
    pages = max(1, doc.pageCount())
    for i in range(pages):
        if i > 0:
            writer.newPage()
        painter.fillRect(QRectF(0, 0, page_w, page_h), QColor(main_bg))
        painter.save()
        painter.translate(0, -i * page_h)
        doc.drawContents(painter, QRectF(0, i * page_h, page_w, page_h))
        painter.restore()
    painter.end()


# ─── XLSX ────────────────────────────────────────────────────────────────────

# Ширина столбца XLSX в символах ≈ (px - 5) / 7. 400 px ≈ 56 символов.
_XLSX_MAX_WIDTH = 56


def export_xlsx(tree, opts, path):
    # Импорт openpyxl изолирован: остальные форматы не требуют этой зависимости.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    cols, header = _table_columns(opts)
    rows = []

    def walk(node, prefix):
        name = node["name"]
        if node["type"] == "account":
            rows.append(_table_row(node, prefix + [name], cols, opts))
        for child in node.get("children", []):
            walk(child, prefix + [name])

    for node in tree:
        walk(node, [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Хранилка"
    ws.append(header)
    for row in rows:
        ws.append(row)

    # Только чёрный текст; выравнивание влево, перенос строк. Цвета темы
    # игнорируются (XLSX — для данных, не для оформления).
    black = Font(color="FF000000")
    black_bold = Font(color="FF000000", bold=True)
    align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for r, row_cells in enumerate(ws.iter_rows(), start=1):
        for cell in row_cells:
            cell.alignment = align
            cell.font = black_bold if r == 1 else black

    # Ширина столбца — по самому длинному фрагменту текста, но не более 400 px.
    for c, title in enumerate(header, start=1):
        longest = len(str(title))
        for row in rows:
            value = row[c - 1] if c - 1 < len(row) else ""
            for line in str(value).split("\n"):
                longest = max(longest, len(line))
        letter = ws.cell(row=1, column=c).column_letter
        ws.column_dimensions[letter].width = min(longest + 2, _XLSX_MAX_WIDTH)

    wb.save(path)


# Соответствие формата → (функция, расширение, фильтр QFileDialog).
FORMATS = {
    "txt": (export_txt, ".txt", "Текстовый файл (*.txt)"),
    "csv": (export_csv, ".csv", "CSV таблица (*.csv)"),
    "xlsx": (export_xlsx, ".xlsx", "Книга Excel (*.xlsx)"),
    "html": (export_html, ".html", "HTML документ (*.html)"),
    "pdf": (export_pdf, ".pdf", "PDF документ (*.pdf)"),
}
