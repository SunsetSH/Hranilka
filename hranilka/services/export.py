"""Экспорт данных Хранилки в человекочитаемые/переносимые форматы.

Форматы: TXT (блокнот), CSV (таблица для Excel/переноса), HTML (оформленный
документ с картинками), XLSX (книга Excel). На вход — структура из
Database.export_subtree(): список узлов дерева, где у каждого узла
type=='account' есть ключи 'card' (как load_account) и 'links'.

ВНИМАНИЕ: экспорт сохраняет выбранные данные в ОТКРЫТОМ виде. Включение паролей
и секретов управляется флагами Options.

Модуль не зависит от Qt (XLSX тянет openpyxl, импорт изолирован внутри функции).
"""
import base64
import csv
import os
from datetime import datetime
from html import escape as esc
from pathlib import Path

from hranilka.core.fin_domain import mask_card_number
from hranilka.core.fin_types import FIN_TYPES
from hranilka.core.nodetypes import is_fin_node, is_server_node
from hranilka.core.util import best_effort_wipe


# Префиксы-«иконки» элементов дерева (как в главном окне).
PREFIX = {"folder": "[+] ", "service": "[o] ", "account": "(i) "}

# Группы полей карточки (для галочек «Что включить»).
# BASIC  — вкладки «База» + «Логин и пароль».
# OTHER  — все остальные вкладки (ПД, секр. вопросы, фраза, 2FA, технические).
# Галерея (картинки) выгружается отдельно (include_gallery) — только в HTML.
GROUP_BASIC = [
    "Название аккаунта", "Адрес сайта (URL)", "Дата создания",
    "Логин", "Пароль", "Пароль сменён", "Сменять пароль каждые",
    "Заметки", "Связанные аккаунты",
]
# Метка резервных кодов — в HTML рендерится особо (многоколоночный список).
LABEL_2FA_CODES = "Резервные коды 2FA"

GROUP_OTHER = [
    "Мобильный номер", "Имя", "Фамилия", "Отчество", "Дата рождения", "Адрес",
    "Секретные вопросы", "Фраза восстановления", "ID устройства",
    LABEL_2FA_CODES, "IP адрес", "Браузер", "ОС",
]

CSV_DELIM = ";"  # разделитель для русского Excel


class Options:
    """Параметры экспорта.

    include_basic   — базовые данные, логин и пароль (вкладки «База»/«Логин»).
    include_other   — остальные поля (ПД, секр. вопросы, фраза, 2FA, технические).
    include_gallery — галерея (изображения и описания); только HTML.
    include_fin     — финансовые записи (карты/кошельки).
    include_fin_secrets — критичные секреты фин-записей (CVV/PIN/seed/ключи и
                    прочие secret-поля). ВЫКЛ по умолчанию: секреты не выводятся
                    вовсе, кроме номера карты (маскируется).
    include_servers — VPS-серверы (docs/ТЗ_VPS_Серверы.md §5). Независимая
                    ветка от include_fin — сервер не финансовый инструмент.
    include_server_secrets — пароли пользователей ОС/панелей, приватные
                    SSH-ключи и passphrase. ВЫКЛ по умолчанию (зеркало
                    include_fin_secrets, без исключения для маскирования —
                    у сервера нет аналога номера карты).
    theme — словарь цветов/шрифта для HTML."""

    def __init__(self, include_basic=True, include_other=True,
                 include_gallery=True, include_fin=True,
                 include_fin_secrets=False, include_servers=True,
                 include_server_secrets=False, title="Вся база", theme=None):
        self.include_basic = include_basic
        self.include_other = include_other
        self.include_gallery = include_gallery
        self.include_fin = include_fin
        self.include_fin_secrets = include_fin_secrets
        self.include_servers = include_servers
        self.include_server_secrets = include_server_secrets
        self.title = title
        self.theme = theme or {}

    def _group_enabled(self, label):
        if label in GROUP_BASIC:
            return self.include_basic
        if label in GROUP_OTHER:
            return self.include_other
        return True


# ─── Общий слой: поля карточки ───────────────────────────────────────────────

def _s(v):
    return "" if v is None else str(v)


def _csv_safe(value):
    """Нейтрализует formula injection в табличных форматах (CSV/XLSX).

    Excel/openpyxl трактуют значение, начинающееся с = + - @ (или с управляющего
    символа перед ними), как формулу — вплоть до DDE/WEBSERVICE-вызовов при
    открытии файла. Префиксуем такие значения апострофом: Excel покажет текст
    как есть и формулу не выполнит. На HTML/TXT не влияет (там не вызывается)."""
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        return "'" + value
    return value


def _account_rows(card, links, opts, skip_empty=True):
    """Упорядоченный список (подпись, значение) для одной карточки.

    Поля выключённых групп (basic/other) полностью пропускаются. skip_empty=True
    — пропускать также пустые поля включённых групп (TXT/HTML); False — оставлять
    их пустыми (CSV, фиксированные колонки)."""
    f = card.get("fields", {})
    p = card.get("personal", {})
    rows = []

    def add(label, value):
        if not opts._group_enabled(label):
            return
        value = _s(value)
        if value == "" and skip_empty:
            return
        rows.append((label, value))

    add("Название аккаунта", f.get("account_name"))
    add("Адрес сайта (URL)", f.get("url"))
    add("Дата создания", f.get("creation_date"))
    add("Логин", f.get("login"))
    add("Пароль", f.get("password"))
    add("Пароль сменён", f.get("password_changed_date"))
    iv = f.get("password_change_interval_days")
    add("Сменять пароль каждые", f"{iv} дн." if iv else "")
    add("Заметки", f.get("notes"))
    add("Связанные аккаунты", "\n".join(l["name"] for l in links) if links else "")

    add("Мобильный номер", p.get("mobile_phone"))
    add("Имя", p.get("first_name"))
    add("Фамилия", p.get("last_name"))
    add("Отчество", p.get("middle_name"))
    add("Дата рождения", p.get("birth_date"))
    add("Адрес", p.get("address"))

    q_lines = []
    for q in card.get("questions", []):
        qt = _s(q.get("q")).strip()
        at = _s(q.get("a"))
        if not qt and not at:
            continue
        q_lines.append(f"{qt} — {at}" if qt else at)
    add("Секретные вопросы", "\n".join(q_lines))

    add("Фраза восстановления", card.get("recovery", {}).get("phrase"))
    add("ID устройства", card.get("recovery", {}).get("device_id"))

    codes = [_s(c) for c in card.get("codes", []) if _s(c).strip()]
    add(LABEL_2FA_CODES, "\n".join(codes))

    add("IP адрес", f.get("ip"))
    add("Браузер", f.get("browser"))
    add("ОС", f.get("os"))

    return rows


# ─── Общий слой: финансовые записи (карты/кошельки) ──────────────────────────
#
# Строки строятся ОДНИМ обходом по дескриптору типа (ItemTypeSpec) — единый код
# для TXT/CSV/HTML/XLSX. Секретные поля (FieldSpec.secret, в т.ч. kind 'seed' и
# секретные поля списков) по умолчанию НЕ выводятся вовсе; исключение — номер
# карты (mask='card'), он выводится маскированным. include_fin_secrets=True —
# всё полностью. item_type вне реестра (битые данные) не роняет экспорт.

def _is_fin(node):
    return is_fin_node(node["type"])


def _fin_prefix(fin):
    """Префикс-«иконка» записи из дескриптора типа ('(?) ' — неизвестный тип)."""
    spec = FIN_TYPES.get((fin or {}).get("item_type"))
    return spec.tree_prefix if spec else "(?) "


def _fin_field_value(field, payload, opts):
    """Значение одного поля с учётом секретности. Возвращает None, если поле
    выводить не следует (скрытый секрет); иначе — строку."""
    raw = _s(payload.get(field.key))
    if field.secret and not opts.include_fin_secrets:
        # Секрет скрыт. Единственное исключение — номер карты: маскируем.
        if field.mask == "card":
            return mask_card_number(raw)
        return None
    return raw


def _fin_list_lines(lst, payload, opts):
    """Строки одного повторяемого блока (ListSpec): по элементу «поле: v / …».
    Секретные значения внутри подчиняются флагу include_fin_secrets."""
    items = payload.get(lst.key)
    if not isinstance(items, list):
        return []
    lines = []
    for elem in items:
        if not isinstance(elem, dict):
            continue
        parts = []
        for f in lst.item_fields:
            if f.secret and not opts.include_fin_secrets:
                continue
            v = _s(elem.get(f.key)).strip()
            if v:
                parts.append(f"{f.label}: {v}")
        if parts:
            lines.append(" / ".join(parts))
    return lines


def _fin_rows(fin, links, opts, skip_empty=True):
    """Упорядоченный список (подпись, значение) для одной финансовой записи.

    Неизвестный item_type → имя + предупреждающая строка (экспорт не падает)."""
    fin = fin or {}
    item_type = fin.get("item_type") or ""
    payload = fin.get("payload") or {}
    spec = FIN_TYPES.get(item_type)
    rows = []
    if spec is None:
        rows.append(("Тип записи", item_type))
        rows.append(("Внимание", "Неизвестный тип записи — поля не распознаны"))
        return rows

    rows.append(("Тип записи", spec.title))
    for field in spec.fields:
        value = _fin_field_value(field, payload, opts)
        if value is None:                 # скрытый секрет — поля нет вовсе
            continue
        if value == "" and skip_empty:
            continue
        rows.append((field.label, value))

    for lst in spec.lists:
        lines = _fin_list_lines(lst, payload, opts)
        if lines:
            rows.append((lst.label, "\n".join(
                f"{i}. {ln}" for i, ln in enumerate(lines, 1))))

    if links:
        rows.append(("Связан с аккаунтами",
                     "\n".join(l["name"] for l in links)))
    return rows


# ─── Общий слой: VPS-серверы ──────────────────────────────────────────────────
#
# Полностью независимая от фин-слоя ветка (docs/ТЗ_VPS_Серверы.md §2/§5):
# сервер — не финансовый инструмент, свой предикат is_server_node, свои
# таблицы описаний полей. Секретные поля (пароли пользователей ОС/панелей,
# приватный SSH-ключ, passphrase) живут только внутри списковых секций
# (os_users/ssh_keys/panels) — у скалярных полей payload секретов нет.

# Префикс листа-сервера в дереве (как в ui/tree.py PREFIX/ui/widgets/srv_linked.py
# SERVER_TREE_PREFIX — модуль экспорта Qt-независим, поэтому константа не
# импортируется оттуда, а дублируется намеренно).
SERVER_PREFIX = "[#] "

# Скалярные поля вкладки «База» карточки сервера (ключ, подпись) — порядок и
# подписи как в ui/server_tabs.py (_SCALAR_KEYS/heading_label), без секретов.
# extra_ips сюда не входит — списковое поле (список строк), своя обработка в
# _server_rows (УИ §2026-07-15, зеркало LABEL_2FA_CODES у аккаунта).
SERVER_SCALAR_FIELDS = (
    ("hosting", "Провайдер / тариф"),
    ("host", "Хост / IP"),
    ("ssh_port", "Порт SSH"),
    ("os_name", "ОС"),
    ("location", "Локация"),
    ("paid_until", "Оплачен до"),
    ("price", "Стоимость"),
    ("notes", "Заметки"),
)

LABEL_EXTRA_IPS = "Доп. IP"

# Списковые секции: (ключ payload, подпись секции, поля [(ключ, подпись, secret)]).
# Ключи и подписи — зеркало ui/server_tabs.py (_OS_USER_FIELDS/_SSH_KEY_FIELDS/
# _PANEL_FIELDS); секретность полей должна совпадать с FieldSpec.secret там же.
SERVER_LIST_SPECS = (
    ("os_users", "Пользователи ОС", (
        ("login", "Логин", False),
        ("password", "Пароль", True),
        ("role", "Роль", False),
        ("label", "Метка", False),
    )),
    ("ssh_keys", "SSH-ключи", (
        ("label", "Метка", False),
        ("key_type", "Тип ключа", False),
        ("public_key", "Публичный ключ", False),
        ("private_key", "Приватный ключ", True),
        ("passphrase", "Passphrase", True),
    )),
    ("panels", "Панели управления", (
        ("panel_type", "Тип панели", False),
        ("url", "URL", False),
        ("login", "Логин", False),
        ("password", "Пароль", True),
    )),
)


def _server_list_lines(items, fields, opts):
    """Строки одной списковой секции сервера: по элементу «поле: v / …».
    Секретные поля подчиняются include_server_secrets (зеркало _fin_list_lines)."""
    if not isinstance(items, list):
        return []
    lines = []
    for elem in items:
        if not isinstance(elem, dict):
            continue
        parts = []
        for key, label, secret in fields:
            if secret and not opts.include_server_secrets:
                continue
            v = _s(elem.get(key)).strip()
            if v:
                parts.append(f"{label}: {v}")
        if parts:
            lines.append(" / ".join(parts))
    return lines


def _server_rows(server, links, opts, skip_empty=True):
    """Упорядоченный список (подпись, значение) для одного сервера."""
    server = server or {}
    payload = server.get("payload") or {}
    rows = [("Тип записи", "VPS-сервер")]

    for key, label in SERVER_SCALAR_FIELDS:
        value = _s(payload.get(key))
        if value == "" and skip_empty:
            continue
        rows.append((label, value))

    ips = payload.get("extra_ips")
    ips_value = "\n".join(_s(ip) for ip in ips if _s(ip).strip()) \
        if isinstance(ips, list) else _s(ips)
    if not (ips_value == "" and skip_empty):
        rows.append((LABEL_EXTRA_IPS, ips_value))

    for list_key, section_label, fields in SERVER_LIST_SPECS:
        lines = _server_list_lines(payload.get(list_key), fields, opts)
        if lines:
            rows.append((section_label, "\n".join(
                f"{i}. {ln}" for i, ln in enumerate(lines, 1))))

    if links:
        rows.append(("Привязан к аккаунтам",
                     "\n".join(l["name"] for l in links)))
    return rows


# ─── Картинки ────────────────────────────────────────────────────────────────

def _img_ext(data):
    if data[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if data[:4] == b"\x89PNG":
        return ".png"
    if data[:3] == b"GIF":
        return ".gif"
    if data[:2] == b"BM":
        return ".bmp"
    return ".png"


def _img_mime(data):
    return {".jpg": "image/jpeg", ".png": "image/png",
            ".gif": "image/gif", ".bmp": "image/bmp"}[_img_ext(data)]


def _has_image(g):
    """Есть ли у элемента галереи картинка — байты уже в снимке (обычный
    режим) либо только id строки БД без байт (bulk.py gallery_ids_only=True,
    потоковый HTML, второй проход M-04: снимок лёгкий, BLOB читается по
    одному через image_provider)."""
    return g.get("data") is not None or g.get("image_id") is not None


def _gallery(node):
    """Картинки аккаунта с непустыми данными."""
    return [g for g in node["card"].get("gallery", []) if _has_image(g)]


def _fin_gallery(node):
    """Картинки финансовой записи с непустыми данными."""
    fin = node.get("fin") or {}
    return [g for g in fin.get("gallery", []) if _has_image(g)]


def _server_gallery(node):
    """Картинки сервера с непустыми данными (зеркало _fin_gallery)."""
    server = node.get("server") or {}
    return [g for g in server.get("gallery", []) if _has_image(g)]


# ─── TXT ─────────────────────────────────────────────────────────────────────

def export_txt(tree, opts, path):
    path = Path(path)
    bar = "=" * 60
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [bar, " ХРАНИЛКА — ЭКСПОРТ", f" {opts.title}    {now}", bar, ""]

    def rows_block(rows, indent, gallery):
        for label, value in rows:
            if "\n" in value:
                lines.append(f"{indent}{label}:")
                for vl in value.split("\n"):
                    lines.append(f"{indent}    {vl}")
            else:
                lines.append(f"{indent}{label}: {value}")
        # Картинки в текст не помещаются — отмечаем только их наличие.
        if gallery:
            lines.append(f"{indent}[Галерея: {len(gallery)} изобр. — доступно в HTML]")
        lines.append(f"{indent}{'-' * 40}")

    def account(node, indent):
        gal = _gallery(node) if opts.include_other else []
        rows_block(_account_rows(node["card"], node.get("links"), opts), indent, gal)

    def fin(node, indent):
        gal = _fin_gallery(node) if opts.include_other else []
        rows_block(_fin_rows(node.get("fin"), node.get("links"), opts),
                   indent, gal)

    def server(node, indent):
        gal = _server_gallery(node) if opts.include_other else []
        rows_block(_server_rows(node.get("server"), node.get("links"), opts),
                   indent, gal)

    def walk(node, depth):
        is_srv = is_server_node(node["type"])
        if is_srv and not opts.include_servers:
            return
        if _is_fin(node) and not opts.include_fin:
            return
        indent = "    " * depth
        if is_srv:
            prefix = SERVER_PREFIX
        elif _is_fin(node):
            prefix = _fin_prefix(node.get("fin"))
        else:
            prefix = PREFIX[node["type"]]
        lines.append(f"{indent}{prefix}{node['name']}")
        if node["type"] == "account":
            account(node, indent + "    ")
        elif _is_fin(node):
            fin(node, indent + "    ")
        elif is_srv:
            server(node, indent + "    ")
        for child in node.get("children", []):
            walk(child, depth + 1)

    for node in tree:
        walk(node, 0)

    path.write_text("\n".join(lines), encoding="utf-8")


# ─── CSV ─────────────────────────────────────────────────────────────────────

def export_csv(tree, opts, path):
    # Колонки зависят от включённых групп. Колонка «Картинок» (количество
    # вложений) относится к «остальным полям» — выводится только с include_other.
    # Финансовые типы имеют иную схему, поэтому пишем вторую таблицу с собственным
    # заголовком ниже первой, а не переменные пары «поле/значение» под заголовок
    # аккаунтов.
    cols, header = _table_columns(opts)
    account_rows = []
    fin_entries = []
    server_entries = []

    def walk(node, prefix):
        name = node["name"]
        if node["type"] == "account":
            account_rows.append(_table_row(node, prefix + [name], cols, opts))
        elif _is_fin(node):
            if opts.include_fin:
                fin_entries.append((node, prefix + [name]))
        elif is_server_node(node["type"]):
            if opts.include_servers:
                server_entries.append((node, prefix + [name]))
        for child in node.get("children", []):
            walk(child, prefix + [name])

    for node in tree:
        walk(node, [])

    def write_entries_table(w, entries, rows_fn):
        labels = []
        for node, _path_parts in entries:
            for label, _value in rows_fn(node):
                if label not in labels:
                    labels.append(label)
        w.writerow([])
        w.writerow(["Путь", "Название"] + labels)
        for node, path_parts in entries:
            values = dict(rows_fn(node))
            w.writerow(
                [_csv_safe(" / ".join(path_parts)), _csv_safe(_s(node["name"]))]
                + [_csv_safe(_s(values.get(label, ""))) for label in labels])

    # utf-8-sig (BOM) — чтобы Excel корректно открыл кириллицу.
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=CSV_DELIM)
        w.writerow(header)
        w.writerows(account_rows)
        if fin_entries:
            write_entries_table(
                w, fin_entries,
                lambda node: _fin_rows(node.get("fin"), node.get("links"), opts))
        if server_entries:
            write_entries_table(
                w, server_entries,
                lambda node: _server_rows(node.get("server"), node.get("links"), opts))


def _table_columns(opts):
    """Колонки таблицы (CSV/XLSX) по включённым группам. Возвращает (cols, header),
    где cols — поля карточки, header — полная строка заголовка с «Путь»/«Картинок»."""
    cols = []
    if opts.include_basic:
        cols += GROUP_BASIC
    if opts.include_other:
        cols += GROUP_OTHER
    header = ["Путь"] + cols + (["Картинок"] if opts.include_other else [])
    return cols, header


def _table_row(node, path_parts, cols, opts):
    """Строка таблицы для одного аккаунта (CSV/XLSX)."""
    d = dict(_account_rows(node["card"], node.get("links"), opts, skip_empty=False))
    row = [" / ".join(path_parts)] + [d.get(l, "") for l in cols]
    if opts.include_other:
        row.append(str(len(_gallery(node))))
    # Нейтрализуем формулы в каждой ячейке (CSV и XLSX используют эту строку).
    return [_csv_safe(_s(c)) for c in row]


# ─── HTML ────────────────────────────────────────────────────────────────────

# Максимум колонок для резервных кодов 2FA в HTML.
HTML_CODES_MAX_COLS = 3

# Галерея в HTML: карточка «картинка + описание» с тонкой рамкой; текст
# обтекает картинку справа, а ниже её — идёт во всю ширину карточки.
# IMG_MAX — потолок миниатюры (обе стороны), CARD_BASE — базовая ширина
# карточки с текстом: сколько карточек влезет в ряд, решает браузер
# (flex-wrap); карточка без описания сжимается до размеров картинки.
HTML_GALLERY_IMG_MAX_PX = 300
HTML_GALLERY_CARD_BASE_PX = 500


def _codes_cell_html(value):
    """Разметка ячейки резервных кодов: список в CSS-колонках.

    column-count задаёт потолок (HTML_CODES_MAX_COLS), а column-width — по
    самому длинному коду — минимальную ширину колонки: браузер сам уменьшает
    число колонок вплоть до одной, если они не влезают в ширину страницы."""
    codes = value.split("\n")
    width_ch = max(len(c) for c in codes) + 2
    items = "".join(f"<div>{esc(c)}</div>" for c in codes)
    return (f"<div class='codes' style='column-width:{width_ch}ch'>"
            f"{items}</div>")


def _default_img_tag(data):
    src = f"data:{_img_mime(data)};base64," + base64.b64encode(data).decode("ascii")
    return f"<img class='shot' src='{src}'>"


def _img_tag_for(g, kind, image_provider):
    """<img> одного элемента галереи.

    g["data"] уже в снимке (обычный режим — прямой вызов _html_document в
    тестах, снимок без gallery_ids_only) — кодируем сразу. Иначе (снимок
    собран с gallery_ids_only=True, M-04 второй проход: снимок лёгкий, BLOB
    не читались) — запрашиваем байты ПО ОДНОЙ картинке через
    image_provider(kind, image_id) и сразу же теряем на них ссылку после
    того, как построили тег: следующая картинка перезапишет локальную
    переменную, предыдущие байты и их base64-строка освобождаются GC раньше,
    чем будет прочитана следующая — пик памяти не растёт с числом картинок."""
    data = g.get("data")
    if data is None:
        image_id = g.get("image_id")
        if image_id is None or image_provider is None:
            return ""                       # защитно: показать нечего
        data = image_provider(kind, image_id)
        if data is None:
            return ""
    return _default_img_tag(data)


def _html_chunks(tree, opts, image_provider=None):
    """Генератор фрагментов HTML-документа — выдаёт их по мере построения, БЕЗ
    накопления в список и без финального join (M-04 второй проход): каждый
    фрагмент отдаётся вызывающему (export_html — пишет сразу в файл) и
    забывается. image_provider — см. _img_tag_for; None — картинки берутся из
    снимка (g["data"]), как раньше (используется тестами, вызывающими
    _html_document(tree, opts) напрямую с уже загруженными BLOB)."""
    th = opts.theme
    font = th.get("font", "Consolas")
    fs = int(th.get("font_size", 14))
    text = th.get("text_color", "#000000")
    main_bg = th.get("main_bg_color", "#F0F0F0")
    tree_bg = th.get("tree_bg_color", "#FFFFFF")
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    yield "<!DOCTYPE html><html><head><meta charset='utf-8'>"
    yield "<title>Хранилка — экспорт</title><style>"
    yield (f"body{{background:{main_bg};color:{text};"
           f"font-family:'{font}',monospace;font-size:{fs}px;margin:24px;}}")
    yield (f".banner{{border:2px solid #808080;background:{tree_bg};"
           f"padding:12px 16px;margin-bottom:20px;}}")
    yield f".banner h1{{margin:0;font-size:{fs + 8}px;letter-spacing:2px;}}"
    yield f"h1.folder{{font-size:{fs + 6}px;border-bottom:2px solid #808080;margin-top:24px;}}"
    yield f"h2.service{{font-size:{fs + 3}px;margin-left:12px;}}"
    # Блок вложенных элементов папки/сервиса: отступ + левая линия + нижняя
    # отбивка — явная иерархия, чтобы следующий сервис/папка не «прилипал»
    # к содержимому предыдущего.
    yield (".group{margin:6px 0 18px 24px;border-left:2px solid #808080;"
           "padding-left:14px;}")
    yield (f".account{{border:2px inset #808080;background:{tree_bg};"
           f"margin:10px 0 10px 24px;padding:8px 12px;}}")
    yield f".account h3{{margin:0 0 6px 0;font-size:{fs + 2}px;}}"
    yield "table.card{border-collapse:collapse;width:100%;}"
    yield "table.card td{border:1px solid #808080;padding:4px 8px;vertical-align:top;}"
    yield "td.label{font-weight:bold;white-space:nowrap;width:200px;}"
    yield "td.value{white-space:pre-wrap;}"
    yield f".codes{{column-count:{HTML_CODES_MAX_COLS};column-gap:24px;}}"
    yield ".codes div{break-inside:avoid;}"
    # Галерея: flex-wrap раскладывает карточки в несколько колонок по ширине
    # окна; align-items:flex-start — рамка каждой карточки по своему содержимому
    # (не тянется до высоты соседей). Внутри карточки картинка — float:left,
    # поэтому текст идёт справа от неё, а ниже картинки — во всю ширину;
    # flow-root не даёт float «выпасть» из рамки.
    yield (".gallery{margin-top:8px;display:flex;flex-wrap:wrap;"
           "gap:10px;align-items:flex-start;}")
    yield (f".gitem{{display:flow-root;"
           f"border:1px solid #808080;background:{tree_bg};padding:8px;"
           f"box-sizing:border-box;max-width:100%;"
           f"flex:1 1 {HTML_GALLERY_CARD_BASE_PX}px;}}")
    # Без описания карточка сжимается до картинки — пустого места нет.
    yield ".gitem.nocap{flex:0 0 auto;}"
    yield (f"img.shot{{max-width:{HTML_GALLERY_IMG_MAX_PX}px;"
           f"max-height:{HTML_GALLERY_IMG_MAX_PX}px;"
           f"border:2px outset #808080;float:left;margin:0 12px 8px 0;}}")
    yield (".cap{font-style:italic;line-height:1.45;margin:0;"
           "white-space:pre-wrap;overflow-wrap:anywhere;}")
    yield "</style></head><body>"
    yield (f"<div class='banner'><h1>ХРАНИЛКА — ЭКСПОРТ</h1>"
           f"{esc(opts.title)} &nbsp;&nbsp; {now}</div>")

    def gallery_html(gal, kind):
        if not (opts.include_gallery and gal):
            return
        yield "<div class='gallery'>"
        for g in gal:
            desc = esc(_s(g.get("desc")))
            cls = "gitem" if desc else "gitem nocap"
            yield f"<div class='{cls}'>"
            yield _img_tag_for(g, kind, image_provider)
            if desc:
                yield f"<div class='cap'>{desc}</div>"
            yield "</div>"
        yield "</div>"

    def account(node):
        yield "<div class='account'>"
        yield f"<h3>(i) {esc(node['name'])}</h3><table class='card'>"
        for label, value in _account_rows(node["card"], node.get("links"), opts):
            cell = (_codes_cell_html(value) if label == LABEL_2FA_CODES
                    else esc(value))
            yield (f"<tr><td class='label'>{esc(label)}</td>"
                   f"<td class='value'>{cell}</td></tr>")
        yield "</table>"
        yield from gallery_html(_gallery(node), "account")
        yield "</div>"

    def fin(node):
        yield "<div class='account'>"
        yield (f"<h3>{esc(_fin_prefix(node.get('fin')))}{esc(node['name'])}</h3>"
               f"<table class='card'>")
        for label, value in _fin_rows(node.get("fin"), node.get("links"), opts):
            yield (f"<tr><td class='label'>{esc(label)}</td>"
                   f"<td class='value'>{esc(value)}</td></tr>")
        yield "</table>"
        yield from gallery_html(_fin_gallery(node), "fin")
        yield "</div>"

    def server(node):
        yield "<div class='account'>"
        yield (f"<h3>{esc(SERVER_PREFIX)}{esc(node['name'])}</h3>"
               f"<table class='card'>")
        for label, value in _server_rows(node.get("server"), node.get("links"), opts):
            yield (f"<tr><td class='label'>{esc(label)}</td>"
                   f"<td class='value'>{esc(value)}</td></tr>")
        yield "</table>"
        yield from gallery_html(_server_gallery(node), "server")
        yield "</div>"

    def walk(node):
        if is_server_node(node["type"]):
            if opts.include_servers:
                yield from server(node)
            return
        if _is_fin(node):
            if opts.include_fin:
                yield from fin(node)
            return
        t = node["type"]
        if t == "account":
            yield from account(node)    # у аккаунта нет потомков — блок не нужен
            return
        if t == "folder":
            yield f"<h1 class='folder'>[+] {esc(node['name'])}</h1>"
        elif t == "service":
            yield f"<h2 class='service'>[o] {esc(node['name'])}</h2>"
        # Вложенные элементы папки/сервиса — в отдельный блок с отступом и
        # левой линией (визуальная отбивка от следующего узла того же уровня).
        yield "<div class='group'>"
        for child in node.get("children", []):
            yield from walk(child)
        yield "</div>"

    for node in tree:
        yield from walk(node)
    yield "</body></html>"


def _html_document(tree, opts):
    """Полный HTML одной строкой (тесты/редкие небольшие вызовы напрямую с уже
    загруженными BLOB в снимке). Реальный экспорт (export_html) пишет
    фрагменты по мере построения, без этого join (M-04 второй проход)."""
    return "".join(_html_chunks(tree, opts))


def export_html(tree, opts, path, image_provider=None):
    """Пишет HTML ЧАСТЯМИ прямо в файл — без накопления списка фрагментов и
    без финального join (M-04 второй проход). image_provider(kind, image_id)
    -> bytes|None задаётся только когда снимок собран с gallery_ids_only=True
    (ExportDialog, HTML+галерея): тогда байты картинок читаются по одной сразу
    в момент вставки в HTML, а не все разом в снимке — пик памяти не растёт с
    объёмом галереи. Без provider — снимок уже содержит все BLOB (обычный
    вызов, как раньше)."""
    with open(path, "w", encoding="utf-8") as f:
        for chunk in _html_chunks(tree, opts, image_provider):
            f.write(chunk)


# ─── XLSX ────────────────────────────────────────────────────────────────────

# Ширина столбца XLSX в символах ≈ (px - 5) / 7. 400 px ≈ 56 символов.
_XLSX_MAX_WIDTH = 56


def export_xlsx(tree, opts, path):
    # Импорт openpyxl изолирован: остальные форматы не требуют этой зависимости.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    account_cols, account_header = _table_columns(opts)
    account_rows = []
    fin_entries = []
    server_entries = []

    def walk(node, prefix):
        name = node["name"]
        if node["type"] == "account":
            account_rows.append(_table_row(node, prefix + [name], account_cols, opts))
        elif _is_fin(node):
            if opts.include_fin:
                fin_entries.append((node, prefix + [name]))
        elif is_server_node(node["type"]):
            if opts.include_servers:
                server_entries.append((node, prefix + [name]))
        for child in node.get("children", []):
            walk(child, prefix + [name])

    for node in tree:
        walk(node, [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Хранилка"
    ws.append(account_header)
    for row in account_rows:
        ws.append(row)

    def append_entries_table(entries, rows_fn):
        """Самостоятельная таблица ниже предыдущих на том же листе. У неё свой
        заголовок: разные записи имеют разные поля, поэтому подгонять их под
        колонки аккаунтов приводило к неразмеченным ячейкам и кривому
        CSV-like представлению внутри XLSX. Возвращает (header, header_row)."""
        if not entries:
            return [], None
        labels = []
        for node, _path_parts in entries:
            for label, _value in rows_fn(node):
                if label not in labels:
                    labels.append(label)
        header = ["Путь", "Название"] + labels
        rows = []
        for node, path_parts in entries:
            values = dict(rows_fn(node))
            rows.append(
                [_csv_safe(" / ".join(path_parts)), _csv_safe(_s(node["name"]))]
                + [_csv_safe(_s(values.get(label, ""))) for label in labels])
        ws.append([])                         # визуальный отступ между таблицами
        ws.append(header)
        header_row = ws.max_row
        for row in rows:
            ws.append(row)
        return header, header_row

    fin_header, fin_header_row = append_entries_table(
        fin_entries,
        lambda node: _fin_rows(node.get("fin"), node.get("links"), opts))
    server_header, server_header_row = append_entries_table(
        server_entries,
        lambda node: _server_rows(node.get("server"), node.get("links"), opts))

    # Только чёрный текст; выравнивание влево, перенос строк. Цвета темы
    # игнорируются (XLSX — для данных, не для оформления).
    black = Font(color="FF000000")
    black_bold = Font(color="FF000000", bold=True)
    align = Alignment(horizontal="left", vertical="top", wrap_text=True)
    bold_rows = {1, fin_header_row, server_header_row} - {None}
    for r, row_cells in enumerate(ws.iter_rows(), start=1):
        for cell in row_cells:
            cell.alignment = align
            cell.font = black_bold if r in bold_rows else black

    # Ширина столбца — по всем таблицам, но не более 400 px.
    max_columns = max(len(account_header), len(fin_header), len(server_header), 1)
    for c in range(1, max_columns + 1):
        longest = 0
        for row_cells in ws.iter_rows(min_col=c, max_col=c):
            value = row_cells[0].value
            for line in str(value or "").split("\n"):
                longest = max(longest, len(line))
        letter = ws.cell(row=1, column=c).column_letter
        ws.column_dimensions[letter].width = min(longest + 2, _XLSX_MAX_WIDTH)

    wb.save(path)


# ─── Атомарная публикация ─────────────────────────────────────────────────────

def write_atomic(func, tree, opts, path, **kwargs):
    """Форматирует экспорт функцией func (одна из FORMATS) во временный файл
    рядом с целевым и публикует его атомарным переименованием (M-04).

    Обрыв/исключение форматтера (например, при обрыве image_provider на
    большой галерее HTML) не оставляет частично записанный или повреждённый
    целевой файл — target либо появляется целиком, либо не появляется вовсе;
    временный файл в любом случае убирается. func сам ничего не знает об
    атомарности — пишет в переданный ему путь как обычно (контракт FORMATS не
    меняется). **kwargs — дополнительные именованные аргументы форматтеру
    (сейчас единственный случай — image_provider у export_html, M-04 второй
    проход); для TXT/CSV/XLSX не передаются, их сигнатура не меняется.

    Временный файл — частично/полностью сформированный экспорт с секретами в
    открытом виде (пароли, коды, содержимое галереи). При обрыве форматтера он
    убирается через best_effort_wipe, а не голым unlink, иначе plaintext мог бы
    остаться в свободных блоках ФС (M-02, как и pre-migrate/restore-копии БД)."""
    target = Path(path)
    tmp = target.with_name(target.name + f".tmp{os.getpid()}")
    try:
        func(tree, opts, str(tmp), **kwargs)
        os.replace(tmp, target)
    finally:
        if tmp.exists():
            best_effort_wipe(str(tmp))


# Соответствие формата → (функция, расширение, фильтр QFileDialog).
FORMATS = {
    "txt": (export_txt, ".txt", "Текстовый файл (*.txt)"),
    "csv": (export_csv, ".csv", "CSV таблица (*.csv)"),
    "xlsx": (export_xlsx, ".xlsx", "Книга Excel (*.xlsx)"),
    "html": (export_html, ".html", "HTML документ (*.html)"),
}
